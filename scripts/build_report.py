from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "COURSE_DESIGN_REPORT.docx"


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = "D9D9D9") -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_text(cell, text: str, bold: bool = False, color: str = "172033") -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_borders(cell)


def add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True, color="FFFFFF")
        set_cell_shading(table.rows[0].cells[index], "2D6CDF")
    for row_index, row in enumerate(rows, start=1):
        cells = table.add_row().cells
        for column_index, value in enumerate(row):
            set_cell_text(cells[column_index], value)
            if row_index % 2 == 0:
                set_cell_shading(cells[column_index], "F4F7FB")
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run(text)


def add_number(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run(text)


def add_field(paragraph, field: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = field
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instruction)
    run._r.append(separate)
    run._r.append(end)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(23, 32, 51)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size in (("Title", 24), ("Heading 1", 15), ("Heading 2", 12)):
        style = styles[style_name]
        style.font.name = "Aptos Display" if style_name == "Title" else "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(5)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Research Paper Information Management System  |  ")
    footer_run.font.name = "Aptos"
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(100, 116, 139)
    add_field(footer, "PAGE")


def build_report() -> None:
    document = Document()
    configure_document(document)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Research Paper Information Management System")

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(18)
    run = subtitle.add_run("Course Design Report and Defense Baseline")
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(45, 108, 223)
    run.bold = True

    metadata = document.add_paragraph()
    metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metadata.add_run("Technology: Python, FastAPI, Jinja2, SQLAlchemy, PostgreSQL\n").bold = True
    metadata.add_run("Application type: Cross-platform web application\n")
    metadata.add_run("Version: Defense baseline")
    document.add_page_break()

    document.add_heading("1 Introduction", level=1)
    document.add_paragraph(
        "This project implements a web-based research paper information management "
        "system. The system stores paper ID, title, authors, institution, "
        "publication, publication date, abstract, and keywords. It provides "
        "authenticated access and the complete record-management workflow required "
        "by the course design brief."
    )
    document.add_paragraph(
        "The current version is the defense baseline. It focuses on reliable CRUD "
        "operations, clear validation, database persistence, file exchange, and a "
        "maintainable object-oriented structure. A later final-defense version can "
        "extend this foundation with a novel intelligent method."
    )

    document.add_heading("2 Objectives", level=1)
    for text in (
        "Provide a structured repository for research paper metadata.",
        "Support secure username and password login.",
        "Provide add, edit, delete, display, search, sort, and counting operations.",
        "Import and export records through a text-compatible CSV format.",
        "Demonstrate object-oriented programming through models and service classes.",
        "Provide a usable responsive web interface with clear primary navigation.",
    ):
        add_bullet(document, text)

    document.add_heading("3 Functional Requirements", level=1)
    for text in (
        "Login with username and password.",
        "Add a paper record.",
        "Delete one or more records.",
        "Modify a selected record.",
        "Display all records.",
        "Search by paper title or author.",
        "Sort by publication date.",
        "Count records by publication.",
        "Import paper information from a text file.",
        "Export all paper information to a text file.",
    ):
        add_number(document, text)

    document.add_heading("4 Technology Selection", level=1)
    document.add_paragraph(
        "The application is implemented with Python, FastAPI, and Jinja2. "
        "FastAPI provides the web application and HTTP routing, while Jinja2 "
        "renders server-side HTML templates. SQLAlchemy provides an "
        "object-relational mapping layer, allowing database tables to be "
        "represented by Python classes. PostgreSQL is used for persistent storage "
        "and Psycopg is the PostgreSQL driver."
    )
    document.add_paragraph(
        "The project uses a virtual environment and stores secrets in an "
        "environment file. The system can be opened from a browser on Windows, "
        "macOS, or Linux. Uvicorn runs the local development server and can later "
        "be placed behind a production reverse proxy."
    )

    document.add_heading("5 System Architecture", level=1)
    add_table(
        document,
        ["Component", "Responsibility"],
        [
            ["models.py", "Defines the User and Paper ORM entities."],
            ["database.py", "Creates the SQLAlchemy engine and transaction scope."],
            ["auth_service.py", "Handles password hashing and authentication."],
            [
                "paper_service.py",
                "Handles validation, CRUD, search, sorting, statistics, import, and export.",
            ],
            [
                "web/app.py",
                "Creates the FastAPI application, sessions, templates, and static files.",
            ],
            [
                "web/routes.py",
                "Provides authentication, CRUD, search, sorting, statistics, import, "
                "and export routes.",
            ],
            ["web/templates/", "Contains the responsive HTML pages."],
            ["web/static/styles.css", "Contains the visual system for the browser interface."],
            ["app.py", "Initializes the database and starts the web server."],
        ],
    )
    document.add_paragraph(
        "This separation keeps route handling focused on HTTP interaction and "
        "keeps business rules testable independently."
    )

    document.add_heading("6 Database Design", level=1)
    document.add_heading("6.1 Users Table", level=2)
    add_table(
        document,
        ["Field", "Type", "Constraint"],
        [
            ["id", "integer", "primary key"],
            ["username", "varchar", "unique, required"],
            ["password_hash", "varchar", "required"],
            ["created_at", "timestamp", "required"],
        ],
    )
    document.add_heading("6.2 Papers Table", level=2)
    add_table(
        document,
        ["Field", "Type", "Constraint"],
        [
            ["id", "integer", "primary key"],
            ["paper_id", "varchar", "unique, required"],
            ["title", "varchar", "required"],
            ["authors", "varchar", "required"],
            ["institution", "varchar", "required"],
            ["publication", "varchar", "required"],
            ["publication_date", "date", "required"],
            ["abstract", "text", "required"],
            ["keywords", "varchar", "required"],
            ["created_at", "timestamp", "required"],
            ["updated_at", "timestamp", "required"],
        ],
    )
    document.add_paragraph(
        "Indexes are added to paper ID, title, authors, publication, and publication "
        "date to support common management and search operations."
    )

    document.add_heading("7 Data Validation", level=1)
    document.add_paragraph(
        "The service layer trims text input and checks every required field. "
        "Publication dates must use YYYY-MM-DD and cannot be in the future. Paper "
        "IDs are unique in the database. Invalid records are rejected with a "
        "user-readable message before the dialog closes."
    )
    document.add_paragraph(
        "User passwords are hashed with Argon2. The application never stores the "
        "original password as a database value."
    )

    document.add_heading("8 Web Interface Design", level=1)
    for text in (
        "Login page.",
        "Library, Statistics, Import, and Export navigation.",
        "Searchable research paper table.",
        "Add, edit, and delete actions.",
        "Expandable abstract previews.",
        "Visible record count and sort direction.",
        "Responsive layouts for desktop and mobile browser widths.",
    ):
        add_bullet(document, text)
    document.add_paragraph(
        "The interface uses a restrained blue and neutral palette, readable "
        "data-dense tables, confirmation before deletion, and a browser-native "
        "date selector."
    )

    document.add_heading("9 File Format", level=1)
    document.add_paragraph(
        "Import and export use UTF-8 CSV, which is a text file format supported by "
        "common spreadsheet and text editors. The header is:"
    )
    code = document.add_paragraph()
    code.paragraph_format.left_indent = Inches(0.25)
    code.paragraph_format.space_after = Pt(8)
    code_run = code.add_run(
        "paper_id,title,authors,institution,publication,publication_date,abstract,keywords"
    )
    code_run.font.name = "Consolas"
    code_run.font.size = Pt(9)
    document.add_paragraph(
        "Import validates each row independently and reports skipped rows while "
        "keeping valid rows. Export writes all records in chronological order."
    )

    document.add_heading("10 Testing", level=1)
    for text in (
        "Valid date conversion.",
        "Rejection of missing required fields.",
        "Rejection of future publication dates.",
        "Python compilation check.",
        "Ruff code-quality check.",
    ):
        add_bullet(document, text)
    document.add_paragraph(
        "The database-backed workflow is tested after the PostgreSQL password is "
        "configured in the .env file."
    )

    document.add_heading("11 Requirement Traceability", level=1)
    document.add_paragraph(
        "The file docs/DEFENSE_CHECKLIST.md maps each course requirement to its "
        "source implementation and suggested defense demonstration."
    )

    document.add_heading("12 Future Work and Novelty Direction", level=1)
    document.add_paragraph(
        "The baseline system is intentionally designed so intelligent features can "
        "be added without rewriting the CRUD foundation. Possible final-defense "
        "research directions include semantic search with document embeddings, "
        "automatic topic clustering, duplicate-paper detection, or paper "
        "recommendation based on abstract similarity."
    )
    document.add_paragraph("One possible future title is:")
    future = document.add_paragraph()
    future.alignment = WD_ALIGN_PARAGRAPH.CENTER
    future_run = future.add_run(
        "Intelligent Research Paper Information Management System Using "
        "Semantic Search and Topic Clustering"
    )
    future_run.bold = True
    future_run.font.color.rgb = RGBColor(45, 108, 223)
    document.add_paragraph(
        "The final method should be selected with the supervisor after the baseline "
        "has been evaluated."
    )

    document.add_heading("13 Conclusion", level=1)
    document.add_paragraph(
        "The project provides a complete starting implementation for a research "
        "web-based research paper information management system. It meets the required management "
        "operations, uses an object-oriented architecture, persists data in "
        "PostgreSQL, validates input, and provides a browser-based foundation "
        "for the final-defense extension."
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    sys.path.insert(0, str(ROOT))
    build_report()
